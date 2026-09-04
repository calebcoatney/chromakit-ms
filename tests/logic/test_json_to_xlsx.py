"""Tests for logic/json_to_xlsx.py — result-JSON to workbook export."""
import datetime
import json
import os
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from logic import json_to_xlsx as jx


# --- fixtures -------------------------------------------------------------

def _peak(compound_id, peak_number, rt, area):
    """A peak dict shaped like logic/json_exporter.py's _serialize_peak output."""
    return {
        "compound_id": compound_id,
        "peak_number": peak_number,
        "retention_time": rt,
        "integrator": "py",
        "width": 0.12,
        "area": area,
        "start_time": rt - 0.05,
        "end_time": rt + 0.05,
        "is_shoulder": False,
        "mol_percent": None,
        "composition_percent": None,
    }


def _rf_json(sample_id, timestamp):
    return {
        "sample_id": sample_id,
        "timestamp": timestamp,
        "method": "TC DEF_GC",
        "detector": "TCD1A",
        "signal": f"Signal: {sample_id}\\TCD1A.ch",
        "notebook": sample_id,
        "peaks": [
            _peak("Hydrogen", 1, 0.374, 385.86),
            _peak("Oxygen", 2, 0.612, 54.27),
            _peak(None, 3, 0.901, 12.5),        # unassigned peak
        ],
        "quantitation": {
            "strategy": "rf_table",
            "rf_unit": "area_per_mol_pct",
            "composition_basis": "mol%",
            "normalized": True,
            "rf_table": [
                {"compound": "Hydrogen", "response_factor": 520.34},
                {"compound": "Oxygen", "response_factor": 54.27},
            ],
        },
    }


def _plain_json(sample_id, timestamp):
    """A non-rf_table result: no quantitation block, no compound_id."""
    data = _rf_json(sample_id, timestamp)
    data.pop("quantitation")
    for peak in data["peaks"]:
        peak.pop("compound_id")
    return data


def _write_run(base, sample_id, data):
    run_dir = os.path.join(str(base), f"{sample_id}.D")
    os.makedirs(run_dir, exist_ok=True)
    with open(os.path.join(run_dir, f"{sample_id}.D - TCD1A.json"), "w") as fh:
        json.dump(data, fh)
    return run_dir


# --- BUG 1: timestamp parsing --------------------------------------------

def test_parse_agilent_timestamp_with_comma():
    # The exact string Agilent ChemStation writes into the result JSON.
    assert jx._parse_timestamp("30-Jul-26, 16:31:10") == \
        datetime.datetime(2026, 7, 30, 16, 31, 10)


def test_parse_agilent_timestamp_variants():
    assert jx._parse_timestamp("30-Jul-2026, 16:31:10") == \
        datetime.datetime(2026, 7, 30, 16, 31, 10)
    assert jx._parse_timestamp("30-Jul-26, 16:31") == \
        datetime.datetime(2026, 7, 30, 16, 31)


def test_parse_timestamp_keeps_legacy_patterns():
    assert jx._parse_timestamp("30-Jul-26 16:31:10") == \
        datetime.datetime(2026, 7, 30, 16, 31, 10)
    assert jx._parse_timestamp("2026-07-30 16:31:10") == \
        datetime.datetime(2026, 7, 30, 16, 31, 10)


def test_parse_timestamp_returns_none_on_junk():
    assert jx._parse_timestamp("not a date") is None
    assert jx._parse_timestamp("") is None


def test_timestamp_cell_is_native_datetime(tmp_path):
    _write_run(tmp_path, "SIG15395", _rf_json("SIG15395.D", "30-Jul-26, 16:31:10"))
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    assert ws.cell(row=2, column=1).value == "Timestamp:"
    assert ws.cell(row=2, column=2).value == datetime.datetime(2026, 7, 30, 16, 31, 10)


def test_runs_sorted_chronologically_regardless_of_walk_order(tmp_path):
    # Written (and walked) out of chronological order on purpose.
    _write_run(tmp_path, "SIG15397", _rf_json("SIG15397.D", "30-Jul-26, 18:05:00"))
    _write_run(tmp_path, "SIG15399", _rf_json("SIG15399.D", "30-Jul-26, 16:31:10"))
    _write_run(tmp_path, "SIG15398", _rf_json("SIG15398.D", "30-Jul-26, 17:12:30"))

    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    sample_ids = [c.value for col in ws.iter_cols(min_col=2, max_col=2)
                  for c in col if isinstance(c.value, str) and c.value.endswith(".D")]
    assert sample_ids == ["SIG15399.D", "SIG15398.D", "SIG15397.D"]


# --- BUG 2: compound IDs and quantitation columns -------------------------

def _header_row(ws):
    for row in ws.iter_rows(min_col=1, max_col=11):
        if row[0].value == "Compound ID":
            return row
    raise AssertionError("no peaks header row found")


def test_rf_json_emits_compound_and_quant_columns(tmp_path):
    _write_run(tmp_path, "SIG15395", _rf_json("SIG15395.D", "30-Jul-26, 16:31:10"))
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    header = _header_row(ws)
    assert [c.value for c in header] == [
        "Compound ID", "Peak #", "Ret Time", "Integrator", "Width", "Area",
        "Start Time", "End Time", "rf", "mol%", "normalized mol%"]
    assert all(c.font.bold for c in header)

    first = header[0].row + 1
    last = first + 2  # three peaks
    assert ws.cell(row=first, column=1).value == "Hydrogen"
    assert ws.cell(row=first, column=9).value == pytest.approx(520.34)
    assert ws.cell(row=first, column=10).value == f"=F{first}/I{first}"
    assert ws.cell(row=first, column=11).value == \
        f"=J{first}/SUM(J{first}:J{last})"

    # Second peak
    assert ws.cell(row=first + 1, column=1).value == "Oxygen"
    assert ws.cell(row=first + 1, column=10).value == f"=F{first+1}/I{first+1}"


def test_unassigned_peak_has_blank_id_and_no_formulas(tmp_path):
    _write_run(tmp_path, "SIG15395", _rf_json("SIG15395.D", "30-Jul-26, 16:31:10"))
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    header = _header_row(ws)
    third = header[0].row + 3  # the compound_id=None peak
    assert ws.cell(row=third, column=1).value in (None, "")
    assert ws.cell(row=third, column=9).value is None    # no rf
    assert ws.cell(row=third, column=10).value is None   # no #DIV/0!
    assert ws.cell(row=third, column=11).value is None


def test_totals_row_after_last_peak(tmp_path):
    _write_run(tmp_path, "SIG15395", _rf_json("SIG15395.D", "30-Jul-26, 16:31:10"))
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    header = _header_row(ws)
    first = header[0].row + 1
    last = first + 2
    totals = last + 1
    assert ws.cell(row=totals, column=10).value == f"=SUM(J{first}:J{last})"
    assert ws.cell(row=totals, column=11).value == f"=SUM(K{first}:K{last})"
    assert ws.cell(row=totals, column=10).font.bold is True
    assert ws.cell(row=totals, column=11).font.bold is True


def test_json_without_quantitation_omits_rf_columns(tmp_path):
    _write_run(tmp_path, "SIG15400", _plain_json("SIG15400.D", "30-Jul-26, 19:00:00"))
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    header = _header_row(ws)
    assert [c.value for c in header if c.value is not None] == [
        "Compound ID", "Peak #", "Ret Time", "Integrator", "Width", "Area",
        "Start Time", "End Time"]

    first = header[0].row + 1
    assert ws.cell(row=first, column=1).value in (None, "")   # blank Compound ID
    assert ws.cell(row=first, column=2).value == 1
    assert ws.cell(row=first, column=6).value == pytest.approx(385.86)
    for row in range(first, first + 4):
        for col in (9, 10, 11):
            assert ws.cell(row=row, column=col).value is None


def test_mixed_runs_only_rf_run_gets_quant_columns(tmp_path):
    _write_run(tmp_path, "SIG15395", _rf_json("SIG15395.D", "30-Jul-26, 16:31:10"))
    _write_run(tmp_path, "SIG15400", _plain_json("SIG15400.D", "30-Jul-26, 19:00:00"))
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    headers = [row for row in ws.iter_rows(min_col=1, max_col=11)
               if row[0].value == "Compound ID"]
    assert len(headers) == 2
    assert headers[0][8].value == "rf"        # rf_table run came first (earlier ts)
    assert headers[1][8].value is None        # plain run has no rf column


def test_empty_rf_table_degrades_gracefully(tmp_path):
    data = _rf_json("SIG15401.D", "30-Jul-26, 20:00:00")
    data["quantitation"]["rf_table"] = []
    _write_run(tmp_path, "SIG15401", data)
    out = tmp_path / "out.xlsx"
    jx.process_json_to_excel(str(tmp_path), str(out))

    ws = load_workbook(str(out))["Chromatograms"]
    header = _header_row(ws)
    assert header[8].value is None
    # compound_id still surfaces even without quantitation columns
    assert ws.cell(row=header[0].row + 1, column=1).value == "Hydrogen"


def test_build_rf_lookup_skips_zero_and_bad_entries():
    data = {"quantitation": {"rf_table": [
        {"compound": "Hydrogen", "response_factor": 520.34},
        {"compound": "Zero", "response_factor": 0.0},        # would be #DIV/0!
        {"compound": "Bad", "response_factor": "n/a"},
        {"compound": None, "response_factor": 1.0},
    ]}}
    assert jx._build_rf_lookup(data) == {"Hydrogen": 520.34}


def test_build_rf_lookup_empty_without_quantitation():
    assert jx._build_rf_lookup({}) == {}
    assert jx._build_rf_lookup({"quantitation": {}}) == {}


# --- import must not open a GUI ------------------------------------------

def test_import_does_not_open_gui():
    """Importing the module must not construct a Tk root or pop dialogs."""
    repo_root = os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))))
    code = (
        "import sys, builtins\n"
        "import tkinter\n"
        "def _boom(*a, **k):\n"
        "    raise AssertionError('GUI opened at import time')\n"
        "tkinter.Tk = _boom\n"
        "import tkinter.filedialog as fd\n"
        "fd.askdirectory = _boom\n"
        "fd.asksaveasfilename = _boom\n"
        "import logic.json_to_xlsx as m\n"
        "assert hasattr(m, 'process_json_to_excel')\n"
        "print('OK')\n"
    )
    proc = subprocess.run([sys.executable, "-c", code], cwd=repo_root,
                          capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    assert "OK" in proc.stdout
